# backend.py
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import pandas as pd
import statsmodels.api as sm
import io, base64, matplotlib.pyplot as plt
import seaborn as sns
from wordcloud import WordCloud
import requests
from docx import Document
from docx.shared import Inches
import json
import numpy as np
import os

app = FastAPI()

# Autoriser tout le monde à appeler ton API (ton site Netlify)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Lire la clé YouTube depuis les variables d'environnement (plus sécurisé)
YOUTUBE_API_KEY = os.environ.get("YOUTUBE_API_KEY", "")

def fig_to_base64(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format='png', bbox_inches='tight')
    buf.seek(0)
    return base64.b64encode(buf.read()).decode('utf-8')

@app.post("/estimation")
async def estimation(
    file: UploadFile = File(...),
    methode: str = Form(...),
    variable_dependante: str = Form(...),
    variables_independantes: str = Form(...)
):
    # Lire CSV
    content = await file.read()
    df = pd.read_csv(io.StringIO(content.decode("utf-8")))
    y = df[variable_dependante]
    X = df[[v.strip() for v in variables_independantes.split(",")]]
    X = sm.add_constant(X)

    # Modèle selon méthode
    if methode == "mco":
        model = sm.OLS(y, X).fit()
    elif methode == "logit":
        model = sm.Logit(y, X).fit(disp=0)
    elif methode == "probit":
        model = sm.Probit(y, X).fit(disp=0)
    else:
        model = sm.OLS(y, X).fit()

    # Extraire résultats
    coeffs = model.params.to_dict()
    pvals = model.pvalues.to_dict()
    stderrs = model.bse.to_dict() if hasattr(model, 'bse') else {}

    # Figures
    figures = []
    
    # 1. Heatmap des corrélations
    fig, ax = plt.subplots(figsize=(6,4))
    cols = [variable_dependante] + [v.strip() for v in variables_independantes.split(",")]
    corr = df[cols].corr()
    sns.heatmap(corr, annot=True, fmt=".2f", cmap="coolwarm", ax=ax)
    ax.set_title("Matrice de corrélation")
    figures.append({"title": "Matrice de corrélation", "base64": fig_to_base64(fig)})
    plt.close(fig)

    # 2. QQ-plot des résidus
    fig, ax = plt.subplots(figsize=(6,4))
    sm.qqplot(model.resid, line='s', ax=ax)
    ax.set_title("QQ-plot des résidus")
    figures.append({"title": "QQ-plot", "base64": fig_to_base64(fig)})
    plt.close(fig)

    # 3. Résidus vs valeurs prédites
    fig, ax = plt.subplots(figsize=(6,4))
    ax.scatter(model.fittedvalues, model.resid, alpha=0.6)
    ax.axhline(0, color='red', linestyle='--')
    ax.set_xlabel("Valeurs prédites")
    ax.set_ylabel("Résidus")
    ax.set_title("Résidus vs prédictions")
    figures.append({"title": "Résidus vs prédictions", "base64": fig_to_base64(fig)})
    plt.close(fig)

    return {
        "resultats": {
            "coefficients": coeffs,
            "pvalues": pvals,
            "stderrs": stderrs,
            "r2": model.rsquared if hasattr(model, 'rsquared') else None,
            "r2_adj": model.rsquared_adj if hasattr(model, 'rsquared_adj') else None,
            "nobs": model.nobs
        },
        "figures": figures
    }

@app.post("/rapport")
async def rapport(data: dict):
    hypotheses = data.get("hypotheses", [])
    methode = data.get("methode", "")
    resultats = data.get("resultats", {})
    figures = data.get("figures", [])

    doc = Document()
    doc.add_heading('Rapport d\'analyse économétrique', 0)
    doc.add_paragraph("Généré par Delta Infinity Pain")

    doc.add_heading('1. Hypothèses', level=1)
    for h in hypotheses:
        doc.add_paragraph(f'- {h}', style='List Bullet')

    doc.add_heading(f'2. Modèle {methode.upper()}', level=1)
    doc.add_paragraph(f"Méthode utilisée : {methode.upper()}")

    doc.add_heading('3. Résultats', level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Variable"
    hdr[1].text = "Coef."
    hdr[2].text = "E.t."
    hdr[3].text = "p-value"
    for var, coef in resultats.get("coefficients", {}).items():
        row = table.add_row().cells
        row[0].text = var
        row[1].text = f"{coef:.4f}"
        row[2].text = f"{resultats.get('stderrs', {}).get(var, 0):.4f}"
        row[3].text = f"{resultats.get('pvalues', {}).get(var, 0):.4f}"

    doc.add_paragraph(f"R² = {resultats.get('r2', 0):.4f}, R² ajusté = {resultats.get('r2_adj', 0):.4f}, n = {resultats.get('nobs', 0)}")

    doc.add_heading('4. Annexes graphiques', level=1)
    for fig in figures:
        img_data = base64.b64decode(fig["base64"])
        img_stream = io.BytesIO(img_data)
        doc.add_picture(img_stream, width=Inches(5))
        doc.add_paragraph(f"Figure : {fig['title']}")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=rapport.docx"})

@app.get("/youtube")
async def youtube(query: str):
    if not YOUTUBE_API_KEY:
        return {"videos": []}
    url = "https://www.googleapis.com/youtube/v3/search"
    params = {
        "part": "snippet",
        "q": query + " économétrie cours",
        "type": "video",
        "maxResults": 5,
        "key": YOUTUBE_API_KEY
    }
    response = requests.get(url, params=params)
    data = response.json()
    videos = []
    for item in data.get("items", []):
        videos.append({
            "title": item["snippet"]["title"],
            "url": f"https://www.youtube.com/watch?v={item['id']['videoId']}",
            "thumbnail": item["snippet"]["thumbnails"]["default"]["url"]
        })
    return {"videos": videos}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)